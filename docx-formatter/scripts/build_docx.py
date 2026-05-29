#!/usr/bin/env python3
"""
默认格式 Word 文档构建器 — 由 docx-formatter skill 调用
用法：作为内联脚本参考，或直接 python3 build_docx.py output.docx
实际使用时由 AI 根据章节结构动态生成完整脚本。
"""

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_CN, FONT_EN = '宋体', 'Times New Roman'

def set_font(run, size=Pt(12), bold=False):
    run.font.name = FONT_EN; run.font.size = size; run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rF = rPr.find(qn('w:rFonts'))
    if rF is None: rF = OxmlElement('w:rFonts'); rPr.insert(0, rF)
    for a, v in [('w:eastAsia', FONT_CN), ('w:ascii', FONT_EN), ('w:hAnsi', FONT_EN)]:
        rF.set(qn(a), v)

def body(doc, text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: pf.first_line_indent = Pt(24)
    r = p.add_run(text); set_font(r, Pt(12))

def item(doc, text):
    """无缩进列表项"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT; pf.first_line_indent = Pt(0)
    r = p.add_run(text); set_font(r, Pt(12))

def heading(doc, text, size=Pt(16), level=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_before = Pt(12); pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT; pf.first_line_indent = Pt(0)
    r = p.add_run(text); set_font(r, size, bold=True)
    if level is not None:
        pPr = p._element.get_or_add_pPr()
        ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), str(level))
        pPr.append(ol)

def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text); set_font(r, Pt(22), bold=True)

def fig_caption(doc, number, desc):
    """图标题:（图X：XXX）居中加粗"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(f'（图{number}：{desc}）')
    set_font(r, Pt(12), bold=True)

def table_row(doc, cells, bold_first=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_before = Pt(2); pf.space_after = Pt(2)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT; pf.first_line_indent = Pt(0)
    r1 = p.add_run(cells[0] + '\t')
    set_font(r1, Pt(11), bold=bold_first)
    r2 = p.add_run(cells[1])
    set_font(r2, Pt(11), bold=False)

# ============ 示例用法 ============
if __name__ == '__main__':
    import sys
    doc = Document()
    for s in doc.sections:
        s.page_width = Cm(21.0); s.page_height = Cm(29.7)
        s.top_margin = Cm(2.54); s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.17); s.right_margin = Cm(3.17)

    title(doc, '示例文档标题')
    heading(doc, '一、第一章', Pt(16), level=1)
    heading(doc, '1.1 第一节', Pt(14), level=2)
    body(doc, '这是正文，首行缩进2字符。')
    item(doc, '这是无缩进列表项。')
    heading(doc, '二、第二章', Pt(16), level=1)
    fig_caption(doc, 1, '示例示意图')

    out = sys.argv[1] if len(sys.argv) > 1 else '/tmp/example.docx'
    doc.save(out)
    print(f'[OK] {out}')
