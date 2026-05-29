#!/usr/bin/env python3
"""Validate strict Chinese official-document DOCX formatting."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


EXPECTED = {
    "margins_cm": {"top": 3.7, "bottom": 3.5, "left": 2.8, "right": 2.6},
    "fonts": {
        "title": ("方正小标宋_GBK", "Times New Roman", 22, False, 34, WD_ALIGN_PARAGRAPH.CENTER, 0),
        "h1": ("方正黑体_GBK", "Times New Roman", 16, False, 29, WD_ALIGN_PARAGRAPH.JUSTIFY, 32),
        "h2": ("方正楷体_GBK", "Times New Roman", 16, False, 29, WD_ALIGN_PARAGRAPH.JUSTIFY, 32),
        "h3": ("方正仿宋_GBK", "Times New Roman", 16, True, 29, WD_ALIGN_PARAGRAPH.JUSTIFY, 32),
        "h4": ("方正仿宋_GBK", "Times New Roman", 16, False, 29, WD_ALIGN_PARAGRAPH.JUSTIFY, 32),
        "body": ("方正仿宋_GBK", "Times New Roman", 16, False, 29, WD_ALIGN_PARAGRAPH.JUSTIFY, 32),
    },
}
CN_DIGITS = "一二三四五六七八九十"
H1_RE = re.compile(rf"^[{CN_DIGITS}]+、")
H2_RE = re.compile(rf"^（[{CN_DIGITS}]+）")
H3_RE = re.compile(r"^\d+\.")
H4_RE = re.compile(r"^（\d+）")


def almost(a, b, tolerance=0.03) -> bool:
    if a is None:
        return False
    return abs(float(a) - float(b)) <= tolerance


def pt_value(length) -> float | None:
    if length is None:
        return None
    return length.pt


def run_fonts(run) -> tuple[str | None, str | None, float | None, bool]:
    r_pr = run._element.find(qn("w:rPr"))
    east = west = None
    size = None
    bold = bool(run.bold)
    if r_pr is not None:
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is not None:
            east = r_fonts.get(qn("w:eastAsia"))
            west = r_fonts.get(qn("w:ascii")) or r_fonts.get(qn("w:hAnsi"))
        sz = r_pr.find(qn("w:sz"))
        if sz is not None:
            size = int(sz.get(qn("w:val"))) / 2
        b = r_pr.find(qn("w:b"))
        if b is not None:
            val = b.get(qn("w:val"))
            bold = val not in ("0", "false", "False")
    return east, west, size, bold


def classify(text: str, index: int) -> str:
    if index == 0:
        return "title"
    if H1_RE.match(text):
        return "h1"
    if H2_RE.match(text):
        return "h2"
    if H3_RE.match(text):
        return "h3"
    if H4_RE.match(text):
        return "h4"
    return "body"


def validate(path: Path, verbose: bool = False) -> int:
    doc = Document(path)
    issues: list[str] = []
    section = doc.sections[0]

    margin_checks = {
        "top": section.top_margin.cm,
        "bottom": section.bottom_margin.cm,
        "left": section.left_margin.cm,
        "right": section.right_margin.cm,
    }
    for key, actual in margin_checks.items():
        expected = EXPECTED["margins_cm"][key]
        if not almost(actual, expected):
            issues.append(f"margin {key}: {actual:.2f}cm != {expected:.2f}cm")

    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    for idx, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        level = classify(text, idx)
        expected_east, expected_west, expected_size, expected_bold, expected_line, expected_align, expected_indent = EXPECTED["fonts"][level]
        if not paragraph.runs:
            issues.append(f"paragraph {idx + 1}: no runs")
            continue
        east, west, size, bold = run_fonts(paragraph.runs[0])
        fmt = paragraph.paragraph_format
        prefix = f"paragraph {idx + 1} [{level}] {text[:24]!r}"
        if east != expected_east:
            issues.append(f"{prefix}: east font {east!r} != {expected_east!r}")
        if west != expected_west:
            issues.append(f"{prefix}: west font {west!r} != {expected_west!r}")
        if not almost(size, expected_size, 0.01):
            issues.append(f"{prefix}: size {size!r} != {expected_size}pt")
        if bold != expected_bold:
            issues.append(f"{prefix}: bold {bold!r} != {expected_bold!r}")
        if fmt.alignment != expected_align:
            issues.append(f"{prefix}: alignment {fmt.alignment!r} != {expected_align!r}")
        if fmt.line_spacing_rule != WD_LINE_SPACING.EXACTLY:
            issues.append(f"{prefix}: line spacing rule is not EXACTLY")
        if not almost(pt_value(fmt.line_spacing), expected_line, 0.01):
            issues.append(f"{prefix}: line spacing {pt_value(fmt.line_spacing)!r} != {expected_line}pt")
        if not almost(pt_value(fmt.space_before), 0, 0.01):
            issues.append(f"{prefix}: space_before {pt_value(fmt.space_before)!r} != 0pt")
        if not almost(pt_value(fmt.space_after), 0, 0.01):
            issues.append(f"{prefix}: space_after {pt_value(fmt.space_after)!r} != 0pt")
        if not almost(pt_value(fmt.first_line_indent), expected_indent, 0.01):
            issues.append(f"{prefix}: first_line_indent {pt_value(fmt.first_line_indent)!r} != {expected_indent}pt")
        if verbose:
            print(f"[CHECK] {idx + 1:02d} {level}: {text}")

    footer = section.footer
    footer_text = "".join(p.text for p in footer.paragraphs)
    has_page_field = any("PAGE" in (node.text or "") for p in footer.paragraphs for node in p._p.iter())
    if not has_page_field and not footer_text.strip().isdigit():
        issues.append("footer: PAGE field not found")
    for paragraph in footer.paragraphs:
        if paragraph.runs:
            _, west, size, _ = run_fonts(paragraph.runs[0])
            if west != "Times New Roman":
                issues.append(f"footer: page number west font {west!r} != 'Times New Roman'")
            if not almost(size, 14, 0.01):
                issues.append(f"footer: page number size {size!r} != 14pt")
            break

    if issues:
        print(f"[FAIL] {path}")
        for issue in issues:
            print(f"  - {issue}")
        return 1
    print(f"[OK] format validation passed: {path}")
    print(f"[OK] checked paragraphs: {len(paragraphs)}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate official DOCX formatting")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--verbose", action="store_true")
    args = parser.parse_args()
    sys.exit(validate(args.docx, args.verbose))


if __name__ == "__main__":
    main()
