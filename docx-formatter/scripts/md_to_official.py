#!/usr/bin/env python3
"""Convert structured Markdown to a strict Chinese official-document DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PAGE_MARGINS_CM = {"top": 3.7, "bottom": 3.5, "left": 2.8, "right": 2.6}
FONT_EAST = {
    "title": "方正小标宋_GBK",
    "h1": "方正黑体_GBK",
    "h2": "方正楷体_GBK",
    "h3": "方正仿宋_GBK",
    "h4": "方正仿宋_GBK",
    "body": "方正仿宋_GBK",
}
FONT_WEST = "Times New Roman"
SIZE_PT = {"title": 22, "h1": 16, "h2": 16, "h3": 16, "h4": 16, "body": 16}
BOLD = {"title": False, "h1": False, "h2": False, "h3": True, "h4": False, "body": False}
LINE_PT = {"title": 34, "h1": 29, "h2": 29, "h3": 29, "h4": 29, "body": 29}
FIRST_INDENT_PT = 32

CN_DIGITS = "一二三四五六七八九十"
H1_PREFIX_RE = re.compile(rf"^[{CN_DIGITS}]+[、.．]\s*")
H2_PREFIX_RE = re.compile(rf"^（[{CN_DIGITS}]+）\s*")
H3_PREFIX_RE = re.compile(r"^\d+[.．]\s*")
H4_PREFIX_RE = re.compile(r"^（\d+）\s*|^\(\d+\)\s*")
LIST_PREFIX_RE = re.compile(r"^[-*]\s+")


def int_to_cn(n: int) -> str:
    if n <= 0 or n > 99:
        return str(n)
    if n <= 10:
        return CN_DIGITS[n - 1]
    tens, ones = divmod(n, 10)
    head = "十" if tens == 1 else CN_DIGITS[tens - 1] + "十"
    return head if ones == 0 else head + CN_DIGITS[ones - 1]


def set_run_font(run, level: str) -> None:
    run.font.name = FONT_WEST
    run.font.size = Pt(SIZE_PT[level])
    run.bold = BOLD[level]
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), FONT_EAST[level])
    r_fonts.set(qn("w:ascii"), FONT_WEST)
    r_fonts.set(qn("w:hAnsi"), FONT_WEST)
    r_fonts.set(qn("w:cs"), FONT_WEST)


def set_paragraph_format(paragraph, level: str) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == "title" else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(LINE_PT[level])
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(0 if level == "title" else FIRST_INDENT_PT)


def classify_markdown_line(line: str) -> tuple[str, str] | None:
    text = line.strip()
    if not text:
        return None
    if text.startswith("# "):
        return "title", text[2:].strip()
    if text.startswith("## "):
        return "h1", text[3:].strip()
    if text.startswith("### "):
        return "h2", text[4:].strip()
    if text.startswith("#### "):
        return "h3", text[5:].strip()
    if text.startswith("##### "):
        return "h4", text[6:].strip()
    text = LIST_PREFIX_RE.sub("", text)
    return "body", text


def strip_existing_number(text: str, level: str) -> str:
    if level == "h1":
        return H1_PREFIX_RE.sub("", text).strip()
    if level == "h2":
        return H2_PREFIX_RE.sub("", text).strip()
    if level == "h3":
        return H3_PREFIX_RE.sub("", text).strip()
    if level == "h4":
        return H4_PREFIX_RE.sub("", text).strip()
    return text.strip()


def apply_numbering(items: list[tuple[str, str]], auto_number: bool) -> list[tuple[str, str]]:
    counters = {"h1": 0, "h2": 0, "h3": 0, "h4": 0}
    out: list[tuple[str, str]] = []
    for level, raw_text in items:
        text = raw_text.strip()
        if not auto_number or level not in counters:
            out.append((level, text))
            continue
        text = strip_existing_number(text, level)
        if level == "h1":
            counters["h1"] += 1
            counters["h2"] = counters["h3"] = counters["h4"] = 0
            text = f"{int_to_cn(counters['h1'])}、{text}"
        elif level == "h2":
            counters["h2"] += 1
            counters["h3"] = counters["h4"] = 0
            text = f"（{int_to_cn(counters['h2'])}）{text}"
        elif level == "h3":
            counters["h3"] += 1
            counters["h4"] = 0
            text = f"{counters['h3']}.{text}"
        elif level == "h4":
            counters["h4"] += 1
            text = f"（{counters['h4']}）{text}"
        out.append((level, text))
    return out


def set_page_number_font(run) -> None:
    run.font.name = FONT_WEST
    run.font.size = Pt(14)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_WEST)
    r_fonts.set(qn("w:hAnsi"), FONT_WEST)
    r_fonts.set(qn("w:cs"), FONT_WEST)


def add_page_number(document: Document) -> None:
    section = document.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for child in list(paragraph._p):
        paragraph._p.remove(child)

    begin = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin._element.append(fld_begin)

    instr = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    instr._element.append(instr_text)

    end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end._element.append(fld_end)

    for run in (begin, instr, end):
        set_page_number_font(run)


def convert(input_path: Path, output_path: Path, auto_number: bool = True) -> None:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(PAGE_MARGINS_CM["top"])
    section.bottom_margin = Cm(PAGE_MARGINS_CM["bottom"])
    section.left_margin = Cm(PAGE_MARGINS_CM["left"])
    section.right_margin = Cm(PAGE_MARGINS_CM["right"])

    items: list[tuple[str, str]] = []
    for line in input_path.read_text(encoding="utf-8").splitlines():
        item = classify_markdown_line(line)
        if item and item[1]:
            items.append(item)
    if not items:
        raise ValueError("input markdown contains no content")

    for level, text in apply_numbering(items, auto_number):
        paragraph = document.add_paragraph()
        set_paragraph_format(paragraph, level)
        run = paragraph.add_run(text)
        set_run_font(run, level)

    add_page_number(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(f"[OK] generated: {output_path}")
    print(f"[INFO] auto_number: {auto_number}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Markdown to strict official DOCX")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--no-auto-number", action="store_true")
    args = parser.parse_args()
    convert(args.input, args.output, auto_number=not args.no_auto_number)


if __name__ == "__main__":
    main()
