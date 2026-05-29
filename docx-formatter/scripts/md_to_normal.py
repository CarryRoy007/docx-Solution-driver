#!/usr/bin/env python3
"""Convert structured Markdown to stable Chinese normal-format DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


FONT_EAST = "宋体"
FONT_WEST = "Times New Roman"
BLACK = "000000"
SIZE_PT = {"title": 16, "h1": 16, "h2": 14, "h3": 12, "h4": 12, "h5": 12, "body": 12, "caption": 12}
CN_DIGITS = "一二三四五六七八九十"
IMAGE_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<path>[^)]+)\)\s*$")


def int_to_cn(n: int) -> str:
    if n <= 0 or n > 99:
        return str(n)
    if n <= 10:
        return CN_DIGITS[n - 1]
    tens, ones = divmod(n, 10)
    head = "十" if tens == 1 else CN_DIGITS[tens - 1] + "十"
    return head if ones == 0 else head + CN_DIGITS[ones - 1]


def set_outline_level(paragraph, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def set_run_font(run, level: str, bold: bool = False) -> None:
    run.font.name = FONT_WEST
    run.font.size = Pt(SIZE_PT[level])
    run.font.bold = bold
    run.font.color.rgb = None
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), FONT_EAST)
    r_fonts.set(qn("w:ascii"), FONT_WEST)
    r_fonts.set(qn("w:hAnsi"), FONT_WEST)
    r_fonts.set(qn("w:cs"), FONT_WEST)
    color = r_pr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        r_pr.append(color)
    color.set(qn("w:val"), BLACK)


def set_para_format(paragraph, level: str) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if level == "title":
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.first_line_indent = Pt(0)
    elif level in {"h1", "h2", "h3", "h4", "h5"}:
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.first_line_indent = Pt(0)
    elif level == "caption":
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.first_line_indent = Pt(0)
    else:
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.first_line_indent = Pt(24)


def classify_line(line: str) -> tuple[str, str] | None:
    text = line.strip()
    if not text:
        return None
    image = IMAGE_RE.match(text)
    if image:
        return "image", text
    for marks, level in [
        ("###### ", "h5"),
        ("##### ", "h4"),
        ("#### ", "h3"),
        ("### ", "h2"),
        ("## ", "h1"),
        ("# ", "title"),
    ]:
        if text.startswith(marks):
            return level, text[len(marks):].strip()
    return "body", text


def strip_existing_number(text: str, level: str) -> str:
    patterns = {
        "h1": rf"^[{CN_DIGITS}]+、\s*",
        "h2": r"^\d+\.\d+\s*",
        "h3": r"^\d+\.\d+\.\d+\s*",
        "h4": r"^\d+\.\d+\.\d+\.\d+\s*",
        "h5": r"^（\d+）\s*|^\(\d+\)\s*",
    }
    pattern = patterns.get(level)
    return re.sub(pattern, "", text).strip() if pattern else text.strip()


def number_items(items: list[tuple[str, str]], auto_number: bool) -> list[tuple[str, str]]:
    counts = {"h1": 0, "h2": 0, "h3": 0, "h4": 0, "h5": 0}
    out: list[tuple[str, str]] = []
    for level, text in items:
        if not auto_number or level not in counts:
            out.append((level, text))
            continue
        text = strip_existing_number(text, level)
        if level == "h1":
            counts["h1"] += 1
            counts["h2"] = counts["h3"] = counts["h4"] = counts["h5"] = 0
            text = f"{int_to_cn(counts['h1'])}、{text}"
        elif level == "h2":
            counts["h2"] += 1
            counts["h3"] = counts["h4"] = counts["h5"] = 0
            text = f"{counts['h1']}.{counts['h2']} {text}"
        elif level == "h3":
            counts["h3"] += 1
            counts["h4"] = counts["h5"] = 0
            text = f"{counts['h1']}.{counts['h2']}.{counts['h3']} {text}"
        elif level == "h4":
            counts["h4"] += 1
            counts["h5"] = 0
            text = f"{counts['h1']}.{counts['h2']}.{counts['h3']}.{counts['h4']} {text}"
        elif level == "h5":
            counts["h5"] += 1
            text = f"（{counts['h5']}）{text}"
        out.append((level, text))
    return out


def add_text_paragraph(doc: Document, level: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    set_para_format(paragraph, level)
    if level in {"h1", "h2", "h3", "h4", "h5"}:
        set_outline_level(paragraph, {"h1": 0, "h2": 1, "h3": 2, "h4": 3, "h5": 4}[level])
    run = paragraph.add_run(text)
    set_run_font(run, level, bold=(level in {"title", "h1", "h2", "h3", "h4", "h5"}))


def add_image(doc: Document, markdown_line: str, base_dir: Path, image_index: int) -> int:
    match = IMAGE_RE.match(markdown_line)
    if not match:
        return image_index
    alt = match.group("alt").strip() or "未命名图示"
    image_path = Path(match.group("path").strip())
    if not image_path.is_absolute():
        image_path = base_dir / image_path
    if not image_path.exists():
        raise FileNotFoundError(f"image not found: {image_path}")
    paragraph = doc.add_paragraph()
    set_para_format(paragraph, "caption")
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(5.7))
    caption = doc.add_paragraph()
    set_para_format(caption, "caption")
    cap_run = caption.add_run(f"图{image_index}：{alt}")
    set_run_font(cap_run, "caption", bold=True)
    return image_index + 1


def convert(input_path: Path, output_path: Path, auto_number: bool = True) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    items = [item for line in input_path.read_text(encoding="utf-8").splitlines() if (item := classify_line(line))]
    items = number_items(items, auto_number=auto_number)
    image_index = 1
    for level, text in items:
        if level == "image":
            image_index = add_image(doc, text, input_path.parent, image_index)
        else:
            add_text_paragraph(doc, level, text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"[OK] generated: {output_path}")
    print(f"[INFO] mode: normal")
    print(f"[INFO] auto_number: {auto_number}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Markdown to stable normal DOCX")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--no-auto-number", action="store_true")
    args = parser.parse_args()
    convert(args.input, args.output, auto_number=not args.no_auto_number)


if __name__ == "__main__":
    main()
