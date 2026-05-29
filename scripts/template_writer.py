#!/usr/bin/env python3
"""基于用户模板的 Markdown → DOCX 写入器."""

import json
import os
import sys
import argparse
import re
from pathlib import Path

try:
    import docx
    from docx.shared import Pt, Inches, Cm, Emu, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


FALLBACK_FORMATTER_NORMAL = os.path.expanduser(
    "~/.config/opencode/skills/docx-formatter/scripts/md_to_normal.py")
FALLBACK_FORMATTER_OFFICIAL = os.path.expanduser(
    "~/.config/opencode/skills/docx-formatter/scripts/md_to_official.py")


def load_template_profile(profile_path):
    if not os.path.exists(profile_path):
        return None
    with open(profile_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def markdown_to_sections(markdown_path):
    with open(markdown_path, 'r', encoding='utf-8') as f:
        content = f.read()

    sections = []
    current_section = {"level": 0, "title": "", "content": [], "images": []}

    for line in content.split('\n'):
        heading_match = re.match(r'^(#{1,6})\s+(.+)', line)
        image_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', line)

        if heading_match:
            if current_section["content"] or current_section["title"]:
                sections.append(current_section)
            level = len(heading_match.group(1))
            sections.append(current_section)
            current_section = {
                "level": level,
                "title": heading_match.group(2).strip(),
                "content": [],
                "images": []
            }
        elif image_match:
            current_section["images"].append({
                "alt": image_match.group(1),
                "path": image_match.group(2)
            })
        elif line.strip():
            current_section["content"].append(line)

    if current_section["content"] or current_section["title"]:
        sections.append(current_section)

    return [s for s in sections if s["title"] or s["content"]]


def write_to_template(template_path, profile, markdown_path, output_path):
    if not HAS_DOCX:
        raise RuntimeError("python-docx not installed")

    doc = docx.Document(template_path)
    sections = markdown_to_sections(markdown_path)

    heading_styles = {}
    if profile:
        for level_key, info in profile.get("heading_profile", {}).items():
            lv = int(level_key.replace("level_", "")) if "level_" in level_key else 0
            heading_styles[lv] = {
                "font_size": info.get("avg_font_size"),
                "font_east": info.get("font_east"),
                "bold": info.get("bold"),
                "align": info.get("align")
            }

    body_style = {}
    if profile and profile.get("body_samples"):
        b = profile["body_samples"][0]
        font = b.get("font", {})
        body_style = {
            "font_size": font.get("size_pt", 12),
            "font_east": font.get("font_east", "宋体"),
            "align": b.get("align", "两端对齐"),
            "indent": b.get("indent")
        }

    body_appended = False
    for section in sections:
        if section["level"] > 0:
            level = section["level"] - 1
            style = heading_styles.get(level, {})

            para = doc.add_paragraph()
            run = para.add_run(section["title"])

            font_size = style.get("font_size", 16 - level * 2)
            run.font.size = Pt(font_size) if font_size else Pt(16 - level * 2)

            if style.get("font_east"):
                rpr = run._element.get_or_add_rPr()
                rFonts = rpr.find(qn('w:rFonts'))
                if rFonts is None:
                    rFonts = __import__('lxml.etree', fromlist=['etree']).SubElement(rpr, qn('w:rFonts'))
                rFonts.set(qn('w:eastAsia'), style["font_east"])

            if style.get("bold"):
                run.bold = True

            align_map = {
                '居中': WD_ALIGN_PARAGRAPH.CENTER,
                '左对齐': WD_ALIGN_PARAGRAPH.LEFT,
                '右对齐': WD_ALIGN_PARAGRAPH.RIGHT,
                '两端对齐': WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            if style.get("align") in align_map:
                para.alignment = align_map[style["align"]]

            if level <= 1:
                para.paragraph_format.outline_level = level
        else:
            for content_line in section["content"]:
                para = doc.add_paragraph(content_line)

                font_size = body_style.get("font_size", 12)
                for run in para.runs:
                    run.font.size = Pt(font_size) if font_size else Pt(12)

                if body_style.get("font_east"):
                    for run in para.runs:
                        rpr = run._element.get_or_add_rPr()
                        rFonts = rpr.find(qn('w:rFonts'))
                        if rFonts is None:
                            from lxml import etree
                            rFonts = etree.SubElement(rpr, qn('w:rFonts'))
                        rFonts.set(qn('w:eastAsia'), body_style["font_east"])

                indent = body_style.get("indent")
                if indent:
                    para.paragraph_format.first_line_indent = Pt(indent)

                body_appended = True

        for img in section.get("images", []):
            img_path = img.get("path", "")
            if os.path.exists(img_path):
                para = doc.add_paragraph()
                run = para.add_run()
                run.add_picture(img_path)

                if img.get("alt"):
                    cap_para = doc.add_paragraph()
                    cap_run = cap_para.add_run(img["alt"])
                    cap_run.bold = True
                    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="基于用户模板写入 DOCX")
    parser.add_argument("--template", required=True, help="模板 .docx 路径")
    parser.add_argument("--markdown", required=True, help="内容 Markdown 路径")
    parser.add_argument("--profile", default=None, help="模板分析报告 JSON")
    parser.add_argument("--output", required=True, help="输出 .docx 路径")
    parser.add_argument("--fallback", choices=["normal", "official"], help="模板不足时的回退格式")
    args = parser.parse_args()

    profile = load_template_profile(args.profile) if args.profile else None

    if profile and profile.get("completeness", {}).get("score", 0) >= 50:
        result = write_to_template(args.template, profile, args.markdown, args.output)
        print(f"已按模板样式写入: {result}")
    elif args.fallback == "official" and os.path.exists(FALLBACK_FORMATTER_OFFICIAL):
        import subprocess
        subprocess.run([sys.executable, FALLBACK_FORMATTER_OFFICIAL, args.markdown, args.output], check=True)
        print(f"模板不完整, 已回退到公文格式: {args.output}")
    else:
        import subprocess
        subprocess.run([sys.executable, FALLBACK_FORMATTER_NORMAL, args.markdown, args.output], check=True)
        print(f"模板不完整, 已回退到常规格式: {args.output}")


if __name__ == "__main__":
    main()
